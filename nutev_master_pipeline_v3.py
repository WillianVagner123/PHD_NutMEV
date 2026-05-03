#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
NutEV Master Pipeline v3
========================

Pipeline único para a pesquisa NutEV:
- gera arquitetura de pastas;
- lê taxonomia editável de palavras-chave (PT + EN);
- gera query pack para bases + web;
- busca em bases bibliográficas e na web geral;
- suporta Google Programmable Search/Custom Search JSON API se o usuário já tiver credenciais;
- usa fallback amplo de busca web em HTML quando Google API não estiver disponível;
- rastreia sites oficiais e páginas de resultados;
- baixa documentos públicos (PDF/DOCX/XLSX/PPTX/TXT/HTML/JSON);
- faz OCR em PDFs escaneados e imagens;
- extrai texto, classifica, deduplica e organiza tabelas.

Importante:
- este script baixa apenas documentos públicos acessíveis sem login/paywall;
- não tenta burlar DRM, paywall ou autenticação;
- para Google, use GOOGLE_API_KEY e GOOGLE_CSE_ID se você já tiver acesso;
- se não houver Google API, o script cai para busca web ampla alternativa + crawler oficial.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import os
import re
import sys
import time
import zipfile
from collections import defaultdict
from dataclasses import asdict, dataclass, fields
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple
from urllib.parse import quote, urlencode, urljoin, urlparse, parse_qs, unquote
import xml.etree.ElementTree as ET

import requests

try:
    import pandas as pd
except Exception:
    print("ERRO: instale pandas e openpyxl: pip install pandas openpyxl")
    raise

try:
    from bs4 import BeautifulSoup
except Exception:
    BeautifulSoup = None

try:
    import pdfplumber
except Exception:
    pdfplumber = None

try:
    import pytesseract
except Exception:
    pytesseract = None

try:
    from PIL import Image
except Exception:
    Image = None

try:
    import docx as py_docx
except Exception:
    py_docx = None

try:
    from sklearn.feature_extraction.text import TfidfVectorizer
    from sklearn.metrics.pairwise import cosine_similarity
except Exception:
    TfidfVectorizer = None
    cosine_similarity = None


# ============================================================
# CONFIG
# ============================================================

USER_AGENT = "NutEV-Master-Pipeline/2026-04-27 (public-docs-search; respectful-use)"
REQUEST_TIMEOUT = 30
REQUEST_PAUSE_SECONDS = 0.35
OCR_RESOLUTION = 300
DEFAULT_RETMAX = 30
DEFAULT_WEB_MAX = 12
DEFAULT_DOWNLOAD_MAX = 120
DEFAULT_CRAWL_PAGES = 60
DEFAULT_CRAWL_LINKS_PER_PAGE = 20
LOCAL_MIN_TEXT_FOR_DEDUP = 500
MAX_DOWNLOAD_BYTES = 80_000_000

SUPPORTED_LOCAL_EXT = (
    ".pdf", ".docx", ".txt", ".csv", ".xlsx", ".xls", ".html", ".htm", ".json",
    ".pptx", ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".webp"
)

DOWNLOADABLE_EXT = {".pdf", ".docx", ".xlsx", ".xls", ".csv", ".txt", ".html", ".htm", ".json", ".pptx", ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".webp"}

PAYWALL_HINTS = [
    "login", "signin", "sign-in", "account", "subscribe", "subscription", "paywall", "purchase",
    "/login", "/signin", "/account", "/subscribe"
]

OFFICIAL_LINK_KEYWORDS = [
    "guideline", "guidelines", "guide", "food guide", "dietary", "nutrition", "healthy diet",
    "statement", "consensus", "recommendation", "position", "policy", "report", "manual", "framework",
    "obesity", "diabetes", "hypertension", "cardiovascular", "lifestyle medicine",
    "diretriz", "guia", "alimentar", "nutrição", "saudável", "consenso", "posicionamento",
    "relatório", "manual", "framework", "obesidade", "diabetes", "hipertensão", "cardiovascular",
    "estilo de vida", "culinary", "culinária", "food literacy", "literacia alimentar"
]

DEFAULT_PROJECT_FOLDERS = [
    "00_config",
    "01_querypacks",
    "02_search_hits/raw",
    "02_search_hits/normalized",
    "02_search_hits/rayyan_ready",
    "03_corpus/03A_bibliographic_oa",
    "03_corpus/03B_web_direct_docs",
    "03_corpus/03C_web_landing_pages",
    "03_corpus/03D_official_seed_docs",
    "03_corpus/03E_manual_drop",
    "04_screening",
    "05_extraction",
    "06_tables",
    "07_logs",
    "08_docs",
]

PUBMED_MESH_HINTS = {
    "lifestyle medicine": '"Lifestyle Medicine"[MeSH Terms]',
    "healthy diet": '"Diet, Healthy"[MeSH Terms]',
    "dietary pattern": '"Dietary Patterns"[All Fields]',
    "dietary patterns": '"Dietary Patterns"[All Fields]',
    "mediterranean diet": '"Mediterranean Diet"[MeSH Terms]',
    "dash diet": '"DASH"[All Fields]',
    "health literacy": '"Health Literacy"[MeSH Terms]',
    "cooking": '"Cooking"[MeSH Terms]',
    "patient compliance": '"Patient Compliance"[MeSH Terms]',
    "feasibility": '"Feasibility Studies"[MeSH Terms]',
    "obesity": '"Obesity"[MeSH Terms]',
    "diabetes": '"Diabetes Mellitus, Type 2"[MeSH Terms]',
    "hypertension": '"Hypertension"[MeSH Terms]',
    "dyslipidemia": '"Dyslipidemias"[MeSH Terms]',
    "metabolic syndrome": '"Metabolic Syndrome"[MeSH Terms]',
}


# ============================================================
# DATA CLASSES
# ============================================================

@dataclass
class SearchRecord:
    uid: str
    workstream: str
    query_variant: str
    source: str
    record_kind: str = "metadata"  # metadata | web | official
    title: str = ""
    snippet: str = ""
    abstract: str = ""
    authors: str = ""
    year: str = ""
    language: str = ""
    doi: str = ""
    pmid: str = ""
    pmcid: str = ""
    journal: str = ""
    publisher: str = ""
    document_type: str = ""
    landing_url: str = ""
    pdf_url: str = ""
    oa_url: str = ""
    file_ext_hint: str = ""
    parent_url: str = ""
    is_oa: bool = False
    keywords_matched: str = ""
    relevance_score: float = 0.0
    prisma_decision: str = "PENDING"
    prisma_reason: str = ""
    source_rank: int = 0
    query_used: str = ""
    notes: str = ""
    collected_at_utc: str = ""

@dataclass
class LocalDocumentRecord:
    doc_id: str
    file_path: str
    file_name: str
    file_ext: str
    file_size_kb: float
    modified_time: str
    inferred_workstream: str
    source_bucket: str
    prisma_decision: str
    prisma_reason: str
    language_hint: str
    used_ocr: bool
    text_len: int
    text_hash: str
    dedup_group: str
    duplicate_of: str
    title_guess: str
    matched_domains: str
    matched_conditions: str
    matched_patterns: str
    matched_outcomes: str
    error: str
    raw_excerpt: str

@dataclass
class DownloadManifestRow:
    workstream: str
    bucket: str
    source: str
    title: str
    url: str
    parent_url: str
    saved_path: str
    status: str
    detail: str
    size_bytes: int
    content_type: str
    collected_at_utc: str


# ============================================================
# UTILITÁRIOS
# ============================================================

def now_utc_iso() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")

def log(msg: str) -> None:
    print(f"[{datetime.now().strftime('%H:%M:%S')}] {msg}")

def ensure_folder(path: Path) -> Path:
    path.mkdir(parents=True, exist_ok=True)
    return path

def safe_sleep(seconds: float = REQUEST_PAUSE_SECONDS) -> None:
    time.sleep(seconds)

def normalize_whitespace(text: str) -> str:
    if not text:
        return ""
    text = text.replace("\x00", " ").replace("�", " ")
    text = re.sub(r"[ \t\r\f\v]+", " ", text)
    text = re.sub(r"\n{2,}", "\n", text)
    return text.strip()

def normalize_title(title: str) -> str:
    title = normalize_whitespace(title).lower()
    title = re.sub(r"[^a-z0-9áàâãéèêíìîóòôõúùûçñü ]+", " ", title)
    title = re.sub(r"\s+", " ", title)
    return title.strip()

def slugify(value: str, max_len: int = 100) -> str:
    if not value:
        return "untitled"
    value = value.lower().strip()
    value = re.sub(r"https?://", "", value)
    value = re.sub(r"[^a-z0-9áàâãéèêíìîóòôõúùûçñü]+", "-", value, flags=re.IGNORECASE)
    value = value.strip("-")
    value = re.sub(r"-{2,}", "-", value)
    return value[:max_len] if value else "untitled"

def stable_hash(*parts: str) -> str:
    joined = "||".join(str(p or "") for p in parts)
    return hashlib.sha256(joined.encode("utf-8", errors="ignore")).hexdigest()

def first_nonempty(*values: Any) -> str:
    for value in values:
        if value is None:
            continue
        sval = str(value).strip()
        if sval:
            return sval
    return ""

def safe_dict(value: Any) -> dict:
    return value if isinstance(value, dict) else {}

def safe_list(value: Any) -> list:
    return value if isinstance(value, list) else []

def infer_language_hint(text: str) -> str:
    text_low = (text or "").lower()
    score_pt = sum(x in text_low for x in ["ção", "ões", "alimentar", "saúde", "obesidade", "diretriz"])
    score_en = sum(x in text_low for x in ["healthy", "guideline", "obesity", "adherence", "implementation"])
    if score_pt > score_en:
        return "pt"
    if score_en > score_pt:
        return "en"
    return "mixed"

def session_factory(email: str | None = None) -> requests.Session:
    s = requests.Session()
    headers = {"User-Agent": USER_AGENT}
    if email:
        headers["From"] = email
    s.headers.update(headers)
    return s

def safe_get(session: requests.Session, url: str, params: dict | None = None) -> requests.Response:
    r = session.get(url, params=params, timeout=REQUEST_TIMEOUT)
    r.raise_for_status()
    safe_sleep()
    return r

def safe_post(session: requests.Session, url: str, data: dict | None = None) -> requests.Response:
    r = session.post(url, data=data, timeout=REQUEST_TIMEOUT)
    r.raise_for_status()
    safe_sleep()
    return r

def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))

def write_json(path: Path, payload: Any) -> Path:
    ensure_folder(path.parent)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return path

def write_text(path: Path, content: str) -> Path:
    ensure_folder(path.parent)
    path.write_text(content, encoding="utf-8")
    return path

def row_dicts_to_excel(sheets: Dict[str, pd.DataFrame], out_path: Path) -> None:
    ensure_folder(out_path.parent)
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        for sheet_name, df in sheets.items():
            safe_name = sheet_name[:31] if len(sheet_name) > 31 else sheet_name
            df.to_excel(writer, sheet_name=safe_name, index=False)

def as_dataframe(rows: Sequence[Any], schema_cls=None) -> pd.DataFrame:
    if not rows:
        if schema_cls is None:
            return pd.DataFrame()
        return pd.DataFrame(columns=[f.name for f in fields(schema_cls)])
    if hasattr(rows[0], "__dataclass_fields__"):
        return pd.DataFrame([asdict(x) for x in rows])
    return pd.DataFrame(rows)

def decode_html_entities(text: str) -> str:
    try:
        from html import unescape
        return unescape(text)
    except Exception:
        return text

def compact_free_text_query(query: str, max_terms: int = 14) -> str:
    terms = re.findall(r'[A-Za-zÀ-ÿ0-9][A-Za-zÀ-ÿ0-9\-]+', query or '')
    stop = {
        'and','or','the','of','for','with','in','on','to','a','an',
        'e','ou','de','da','do','das','dos','para','com','em','na','no','nas','nos'
    }
    cleaned = []
    seen = set()
    for term in terms:
        low = term.lower()
        if len(low) < 3 or low in stop:
            continue
        if low not in seen:
            cleaned.append(term)
            seen.add(low)
        if len(cleaned) >= max_terms:
            break
    return ' '.join(cleaned)

def detect_extension_from_response(url: str, content_type: str = "") -> str:
    ctype = (content_type or '').lower()
    path = urlparse(url).path.lower()
    if path.endswith('.pdf') or 'application/pdf' in ctype:
        return '.pdf'
    if path.endswith('.docx') or 'application/vnd.openxmlformats-officedocument.wordprocessingml.document' in ctype:
        return '.docx'
    if path.endswith('.xlsx') or 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet' in ctype:
        return '.xlsx'
    if path.endswith('.xls') or 'application/vnd.ms-excel' in ctype:
        return '.xls'
    if path.endswith('.csv') or 'text/csv' in ctype:
        return '.csv'
    if path.endswith('.txt') or 'text/plain' in ctype:
        return '.txt'
    if path.endswith('.pptx') or 'application/vnd.openxmlformats-officedocument.presentationml.presentation' in ctype:
        return '.pptx'
    if path.endswith('.png') or 'image/png' in ctype:
        return '.png'
    if path.endswith('.jpg') or path.endswith('.jpeg') or 'image/jpeg' in ctype:
        return '.jpg'
    if path.endswith('.tif') or path.endswith('.tiff') or 'image/tiff' in ctype:
        return '.tif'
    if path.endswith('.webp') or 'image/webp' in ctype:
        return '.webp'
    if path.endswith('.json') or 'application/json' in ctype:
        return '.json'
    if path.endswith('.html') or path.endswith('.htm') or 'text/html' in ctype:
        return '.html'
    return Path(path).suffix.lower() or '.bin'

def is_public_downloadable_url(url: str) -> bool:
    low = (url or "").lower()
    if not low.startswith(("http://", "https://")):
        return False
    if any(h in low for h in PAYWALL_HINTS):
        return False
    ext = detect_extension_from_response(low, "")
    if ext in DOWNLOADABLE_EXT:
        return True
    return False

def stream_download(session: requests.Session, url: str, out_path: Path, max_bytes: int = MAX_DOWNLOAD_BYTES) -> tuple[bool, str, int]:
    try:
        with session.get(url, stream=True, timeout=REQUEST_TIMEOUT) as r:
            r.raise_for_status()
            content_type = r.headers.get('Content-Type', '')
            size = 0
            ensure_folder(out_path.parent)
            with out_path.open('wb') as f:
                for chunk in r.iter_content(chunk_size=65536):
                    if not chunk:
                        continue
                    size += len(chunk)
                    if size > max_bytes:
                        return False, f'Arquivo excede limite configurado ({max_bytes} bytes)', size
                    f.write(chunk)
            safe_sleep()
            return True, content_type, size
    except Exception as exc:
        return False, str(exc), 0

def filename_from_record(record: SearchRecord, ext: str, bucket: str = "") -> str:
    year = record.year or "nd"
    source = slugify(record.source, 18)
    title_chunk = slugify(record.title or record.snippet or 'documento', 70)
    bucket_chunk = slugify(bucket or record.record_kind or 'src', 18)
    id_chunk = slugify(record.doi.replace('/', '_'), 30) if record.doi else (record.pmid or record.uid)
    return f"NUTEV__{record.workstream}__{bucket_chunk}__{source}__{year}__{title_chunk}__{id_chunk}{ext}"

def url_to_html_filename(workstream: str, source: str, url: str, title: str = "") -> str:
    host = slugify(urlparse(url).netloc or "site", 30)
    tail = slugify(title or urlparse(url).path or "page", 70)
    return f"NUTEV__{workstream}__landing__{slugify(source,18)}__{host}__{tail}__{stable_hash(url)[:10]}.html"

def is_allowed_domain(url: str, allowed_domains: Sequence[str]) -> bool:
    try:
        host = urlparse(url).netloc.lower()
    except Exception:
        return False
    return any(host == dom.lower() or host.endswith("." + dom.lower()) for dom in allowed_domains)

def looks_like_relevant_link(url: str, text: str) -> bool:
    blob = f"{url} {text}".lower()
    ext = detect_extension_from_response(url, "")
    if ext in DOWNLOADABLE_EXT:
        return True
    return any(k in blob for k in OFFICIAL_LINK_KEYWORDS)

def normalize_url_from_search(href: str) -> str:
    if not href:
        return ""
    if "duckduckgo.com/l/?" in href or "duckduckgo.com/l/?" in href:
        parsed = urlparse(href)
        qs = parse_qs(parsed.query)
        uddg = qs.get("uddg", [])
        if uddg:
            return unquote(uddg[0])
    return href

def html_to_text(html: str) -> str:
    if not html:
        return ""
    if BeautifulSoup is None:
        return re.sub(r"<[^>]+>", " ", html)
    soup = BeautifulSoup(html, "html.parser")
    return soup.get_text("\n", strip=True)

def make_uid(source: str, workstream: str, title: str, year: str, extra: str = "") -> str:
    return stable_hash(source, workstream, normalize_title(title), year, extra)[:16]

def build_default_paths(project_root: Path) -> Dict[str, Path]:
    return {
        "taxonomy": project_root / "00_config" / "nutev_keyword_taxonomy_v3.json",
        "sources": project_root / "00_config" / "nutev_official_sources_manifest_v3.json",
        "query_pack_json": project_root / "01_querypacks" / "nutev_query_pack_v3.json",
        "query_pack_md": project_root / "01_querypacks" / "NUTEV_QUERY_PACK_v3.md",
        "search_master_csv": project_root / "02_search_hits" / "normalized" / "NUTEV_SEARCH_MASTER.csv",
        "search_master_xlsx": project_root / "02_search_hits" / "normalized" / "NUTEV_SEARCH_MASTER.xlsx",
        "rayyan_csv": project_root / "02_search_hits" / "rayyan_ready" / "NUTEV_RAYYAN_READY.csv",
        "download_manifest_csv": project_root / "02_search_hits" / "raw" / "NUTEV_DOWNLOAD_MANIFEST.csv",
        "candidate_links_csv": project_root / "02_search_hits" / "raw" / "NUTEV_CANDIDATE_LINKS.csv",
        "local_xlsx": project_root / "05_extraction" / "NUTEV_LOCAL_ANALYSIS_v3.xlsx",
        "master_tables_xlsx": project_root / "06_tables" / "NUTEV_MASTER_TABLES_v3.xlsx",
        "keywords_md": project_root / "08_docs" / "NUTEV_KEYWORDS_EDITABLE_v3.md",
        "naming_examples": project_root / "08_docs" / "NUTEV_EXEMPLOS_NOMES_DE_ARQUIVO_v3.txt",
        "log_json": project_root / "07_logs" / "run_log_v3.json",
    }

def scaffold_project(project_root: Path) -> None:
    for rel in DEFAULT_PROJECT_FOLDERS:
        ensure_folder(project_root / rel)


# ============================================================
# CONFIG / TAXONOMIA
# ============================================================

def load_taxonomy(taxonomy_path: Path) -> dict:
    if not taxonomy_path.exists():
        raise FileNotFoundError(f"Taxonomia não encontrada: {taxonomy_path}")
    return load_json(taxonomy_path)

def load_seed_manifest(path: Path) -> dict:
    if not path.exists():
        return {"sources": []}
    return load_json(path)

def collect_term_bank(taxonomy: dict, workstream_key: str) -> Dict[str, List[str]]:
    work = taxonomy["workstreams"][workstream_key]
    bank: Dict[str, List[str]] = defaultdict(list)

    for block_name in work.get("focus_blocks", []):
        block = taxonomy["global"].get(block_name, {})
        if isinstance(block, dict):
            for _, terms in block.items():
                bank[block_name].extend(terms)

    bank["population_terms"].extend(work.get("population_terms", []))
    bank["condition_terms"].extend(work.get("condition_terms", []))
    bank["web_query_hints"].extend(work.get("web_query_hints", []))

    for ckey in work.get("clinical_keys", []):
        bank["clinical_terms"].extend(taxonomy.get("clinical", {}).get(ckey, []))

    for dkey in work.get("document_type_keys", []):
        bank["document_type_terms"].extend(taxonomy["global"]["document_types"].get(dkey, []))

    for okey in work.get("priority_outcomes", []):
        bank["outcome_terms"].extend(taxonomy.get("outcomes", {}).get(okey, []))

    for _, terms in taxonomy["global"].get("diet_patterns", {}).items():
        bank["pattern_terms"].extend(terms)
    for _, terms in taxonomy["global"].get("implementation_behavior", {}).items():
        bank["implementation_terms"].extend(terms)
    for _, terms in taxonomy["global"].get("nutrition_domains", {}).items():
        bank["nutrition_terms"].extend(terms)
    for _, terms in taxonomy["global"].get("lifestyle_medicine_pillars", {}).items():
        bank["lifestyle_terms"].extend(terms)

    final_bank = {}
    for key, terms in bank.items():
        seen = set()
        cleaned = []
        for term in terms:
            t = normalize_whitespace(term)
            low = t.lower()
            if t and low not in seen:
                cleaned.append(t)
                seen.add(low)
        final_bank[key] = cleaned
    return final_bank

def render_taxonomy_markdown(taxonomy: dict) -> str:
    lines = ["# NutEV - Banco editável de palavras-chave v3", ""]
    lines.append("Edite este arquivo junto com o JSON quando quiser ampliar, polir ou cortar termos.")
    lines.append("")
    for ws_key, ws in taxonomy["workstreams"].items():
        lines.append(f"## {ws_key}")
        lines.append("")
        lines.append(f"**Título:** {ws['title']}")
        lines.append("")
        lines.append(f"**Pergunta de pesquisa:** {ws['research_question']}")
        lines.append("")
        bank = collect_term_bank(taxonomy, ws_key)
        for section in ["population_terms", "condition_terms", "clinical_terms", "pattern_terms", "nutrition_terms", "implementation_terms", "outcome_terms", "document_type_terms", "web_query_hints"]:
            if bank.get(section):
                lines.append(f"### {section}")
                lines.append("")
                lines.append(", ".join(bank[section]))
                lines.append("")
    return "\n".join(lines)


# ============================================================
# QUERY BUILDERS
# ============================================================

def take_terms(terms: List[str], limit: int) -> List[str]:
    return terms[:limit] if len(terms) > limit else terms

def pubmed_or_block(terms: List[str], field: str = "TIAB") -> str:
    parts = []
    for t in terms:
        low = t.lower()
        if low in PUBMED_MESH_HINTS:
            parts.append(PUBMED_MESH_HINTS[low])
        else:
            escaped = t.replace('"', '\\"')
            parts.append(f'"{escaped}"[{field}]')
    return "(" + " OR ".join(parts) + ")" if parts else ""

def plain_or_block(terms: List[str]) -> str:
    parts = []
    for t in terms:
        escaped = t.replace('"', '\\"')
        if " " in escaped:
            parts.append(f'"{escaped}"')
        else:
            parts.append(escaped)
    return "(" + " OR ".join(parts) + ")" if parts else ""

def build_web_queries(bank: Dict[str, List[str]], workstream_key: str) -> Dict[str, str]:
    lifestyle = take_terms(bank.get("lifestyle_terms", []), 8)
    patterns = take_terms(bank.get("pattern_terms", []), 8)
    clinical = take_terms(bank.get("clinical_terms", []) or bank.get("condition_terms", []), 8)
    outcomes = take_terms(bank.get("outcome_terms", []), 8)
    docs = take_terms(bank.get("document_type_terms", []), 8)
    impl = take_terms(bank.get("implementation_terms", []), 10)
    hints = take_terms(bank.get("web_query_hints", []), 6)

    pt_terms = [t for t in lifestyle + patterns + clinical + impl + docs + hints if any(ch in t for ch in "çãõáéíóú")]
    en_terms = [t for t in lifestyle + patterns + clinical + impl + docs + hints if t not in pt_terms]

    q_balanced_en = " ".join(take_terms(en_terms, 14))
    q_balanced_pt = " ".join(take_terms(pt_terms or bank.get("condition_terms", []), 14))
    q_implementation = " ".join(take_terms(impl + outcomes + patterns, 14))
    q_guidelines = " ".join(take_terms(clinical + docs + hints + patterns, 14))
    q_framework = " ".join(take_terms(lifestyle + impl + bank.get("nutrition_terms", []), 14))

    # queries mais enxutas para não explodir em motores web
    return {
        "balanced_en": compact_free_text_query(q_balanced_en, 12),
        "balanced_pt": compact_free_text_query(q_balanced_pt, 12),
        "implementation": compact_free_text_query(q_implementation, 12),
        "guidelines_docs": compact_free_text_query(q_guidelines, 12),
        "framework": compact_free_text_query(q_framework, 12),
    }

def build_query_variants(taxonomy: dict, workstream_key: str) -> Dict[str, Dict[str, str]]:
    bank = collect_term_bank(taxonomy, workstream_key)

    lifestyle = take_terms(bank.get("lifestyle_terms", []), 12)
    nutrition = take_terms(bank.get("nutrition_terms", []), 16)
    implementation = take_terms(bank.get("implementation_terms", []), 16)
    patterns = take_terms(bank.get("pattern_terms", []), 14)
    clinical = take_terms(bank.get("clinical_terms", []) or bank.get("condition_terms", []), 12)
    population = take_terms(bank.get("population_terms", []), 8)
    document_types = take_terms(bank.get("document_type_terms", []), 14)
    outcomes = take_terms(bank.get("outcome_terms", []), 12)

    pubmed_balanced = " AND ".join(filter(None, [
        pubmed_or_block(lifestyle[:8] + patterns[:6], "TIAB"),
        pubmed_or_block(population + clinical[:8], "TIAB"),
        pubmed_or_block(nutrition[:8] + implementation[:8], "TIAB"),
        pubmed_or_block(document_types[:8], "TIAB"),
    ]))

    pubmed_broad = " AND ".join(filter(None, [
        pubmed_or_block(lifestyle[:6] + patterns[:6], "TIAB"),
        pubmed_or_block(population + clinical[:6], "TIAB"),
        pubmed_or_block(document_types[:6], "TIAB"),
    ]))

    pubmed_implementation = " AND ".join(filter(None, [
        pubmed_or_block(lifestyle[:6] + patterns[:5], "TIAB"),
        pubmed_or_block(implementation[:10], "TIAB"),
        pubmed_or_block(outcomes[:8] + clinical[:6], "TIAB"),
    ]))

    europepmc_balanced = " AND ".join(filter(None, [
        plain_or_block(lifestyle[:8] + patterns[:6]),
        plain_or_block(population + clinical[:8]),
        plain_or_block(nutrition[:8] + implementation[:8]),
        plain_or_block(document_types[:8]),
    ]))

    openalex_balanced = " ".join(
        lifestyle[:4] + patterns[:4] + clinical[:4] + implementation[:4] + document_types[:4]
    ).strip()

    return {
        "pubmed": {
            "broad": pubmed_broad,
            "balanced": pubmed_balanced,
            "implementation": pubmed_implementation,
        },
        "europepmc": {
            "balanced": europepmc_balanced,
            "implementation": " AND ".join(filter(None, [
                plain_or_block(lifestyle[:6] + patterns[:5]),
                plain_or_block(implementation[:10]),
                plain_or_block(outcomes[:8]),
            ])),
        },
        "openalex": {
            "balanced": openalex_balanced,
            "narrow": compact_free_text_query(" ".join(patterns[:6] + clinical[:6] + document_types[:6]), 12),
        },
        "crossref": {
            "balanced": compact_free_text_query(openalex_balanced, 14),
            "narrow": compact_free_text_query(" ".join(clinical[:5] + document_types[:5] + outcomes[:5]), 12),
        },
        "doaj": {
            "balanced": compact_free_text_query(" ".join(patterns[:6] + implementation[:6] + clinical[:6]), 12),
        },
        "web": build_web_queries(bank, workstream_key),
    }

def build_query_pack(taxonomy: dict) -> Dict[str, Any]:
    pack = {"generated_at_utc": now_utc_iso(), "workstreams": {}}
    for workstream_key in taxonomy["workstreams"].keys():
        pack["workstreams"][workstream_key] = {
            "title": taxonomy["workstreams"][workstream_key]["title"],
            "research_question": taxonomy["workstreams"][workstream_key]["research_question"],
            "queries": build_query_variants(taxonomy, workstream_key),
        }
    return pack

def write_query_pack_files(pack: dict, project_root: Path) -> Tuple[Path, Path]:
    paths = build_default_paths(project_root)
    write_json(paths["query_pack_json"], pack)
    md_lines = ["# NutEV Query Pack v3", "", f"Gerado em: {pack['generated_at_utc']}", ""]
    for ws_key, payload in pack["workstreams"].items():
        md_lines.append(f"## {ws_key}")
        md_lines.append("")
        md_lines.append(f"**Título:** {payload['title']}")
        md_lines.append("")
        md_lines.append(f"**Pergunta de pesquisa:** {payload['research_question']}")
        md_lines.append("")
        for source, variants in payload["queries"].items():
            md_lines.append(f"### {source}")
            md_lines.append("")
            for variant_name, query in variants.items():
                md_lines.append(f"**{variant_name}**")
                md_lines.append("")
                md_lines.append("```")
                md_lines.append(query)
                md_lines.append("```")
                md_lines.append("")
    write_text(paths["query_pack_md"], "\n".join(md_lines))
    return paths["query_pack_json"], paths["query_pack_md"]


# ============================================================
# SCORING E DEDUP
# ============================================================

def term_hits(text: str, terms: Sequence[str]) -> List[str]:
    low = (text or "").lower()
    hits = []
    for term in terms:
        t = term.lower()
        if t and t in low:
            hits.append(term)
    seen = set()
    out = []
    for h in hits:
        if h.lower() not in seen:
            out.append(h)
            seen.add(h.lower())
    return out

def score_search_record(record: SearchRecord, taxonomy: dict) -> Tuple[float, List[str], str, str]:
    bank = collect_term_bank(taxonomy, record.workstream)
    text = " ".join([
        record.title, record.snippet, record.abstract, record.document_type,
        record.journal, record.publisher, record.landing_url, record.notes
    ]).lower()

    matched = []
    score = 0.0
    scoring_groups = [
        ("clinical_terms", 2.2),
        ("nutrition_terms", 1.8),
        ("pattern_terms", 1.8),
        ("implementation_terms", 1.6),
        ("outcome_terms", 1.4),
        ("document_type_terms", 1.2),
    ]
    for group_name, weight in scoring_groups:
        hits = term_hits(text, bank.get(group_name, []))
        if hits:
            matched.extend(hits[:8])
            score += weight * min(len(hits), 6)

    if record.doi:
        score += 0.6
    if record.abstract:
        score += 0.8
    if record.is_oa:
        score += 0.5
    if record.pdf_url or record.file_ext_hint in [".pdf", ".docx", ".xlsx", ".pptx"]:
        score += 0.8
    if record.record_kind == "web":
        score += 0.4
    if record.source in taxonomy["workstreams"][record.workstream].get("source_priority", []):
        priority_rank = taxonomy["workstreams"][record.workstream]["source_priority"].index(record.source)
        score += max(0.1, 0.6 - (0.08 * priority_rank))

    title_low = (record.title or "").lower()
    if any(x in title_low for x in ["guideline", "diretriz", "consensus", "statement", "guidance", "guia"]):
        score += 1.3
    if any(x in title_low for x in ["review", "meta-analysis", "revisão", "metanálise"]):
        score += 0.9
    if any(x in title_low for x in ["trial", "ensaio", "randomized", "pilot", "feasibility"]):
        score += 0.8
    if record.file_ext_hint in [".pdf", ".docx", ".xlsx", ".pptx"]:
        score += 0.4

    if score >= 8.0:
        decision, reason = "INCLUDE", "Alta convergência terminológica com o workstream"
    elif score >= 4.8:
        decision, reason = "UNCERTAIN", "Convergência parcial; revisar manualmente"
    else:
        decision, reason = "EXCLUDE", "Baixa aderência ao escopo-alvo"
    return round(score, 2), matched[:20], decision, reason

def preferred_record(a: SearchRecord, b: SearchRecord) -> SearchRecord:
    score_a = (bool(a.abstract), bool(a.doi), a.is_oa, bool(a.pdf_url), bool(a.landing_url), len(a.title), len(a.abstract), a.relevance_score)
    score_b = (bool(b.abstract), bool(b.doi), b.is_oa, bool(b.pdf_url), bool(b.landing_url), len(b.title), len(b.abstract), b.relevance_score)
    return a if score_a >= score_b else b

def deduplicate_records(records: Sequence[SearchRecord]) -> List[SearchRecord]:
    by_key: Dict[str, SearchRecord] = {}
    for record in records:
        if record.doi:
            key = f"doi::{record.doi.lower()}"
        elif record.pmid:
            key = f"pmid::{record.pmid}"
        elif record.landing_url:
            key = f"url::{record.landing_url.lower().rstrip('/')}"
        else:
            key = f"title::{normalize_title(record.title)}::{record.year}"
        if key not in by_key:
            by_key[key] = record
        else:
            by_key[key] = preferred_record(by_key[key], record)

    deduped = list(by_key.values())

    if TfidfVectorizer is not None and cosine_similarity is not None and len(deduped) > 2:
        titles = [normalize_title(x.title) for x in deduped]
        if sum(bool(t) for t in titles) >= 2:
            X = TfidfVectorizer(min_df=1, ngram_range=(1, 2)).fit_transform(titles)
            sim = cosine_similarity(X)
            keep = [True] * len(deduped)
            for i in range(len(deduped)):
                if not keep[i]:
                    continue
                for j in range(i + 1, len(deduped)):
                    if not keep[j]:
                        continue
                    if deduped[i].year and deduped[j].year and deduped[i].year != deduped[j].year:
                        continue
                    if sim[i, j] >= 0.96:
                        best = preferred_record(deduped[i], deduped[j])
                        if best is deduped[i]:
                            keep[j] = False
                        else:
                            keep[i] = False
                            break
            deduped = [x for x, k in zip(deduped, keep) if k]
    return deduped


# ============================================================
# SEARCH APIS - BIBLIOGRÁFICAS
# ============================================================

def decode_openalex_abstract(value: Any) -> str:
    if not isinstance(value, dict):
        return ""
    pairs = []
    for word, positions in value.items():
        for pos in positions:
            pairs.append((pos, word))
    if not pairs:
        return ""
    pairs.sort(key=lambda x: x[0])
    return " ".join(word for _, word in pairs)

def search_pubmed(session: requests.Session, workstream: str, query_variant: str, query: str, email: str, retmax: int) -> List[SearchRecord]:
    esearch_url = "https://eutils.ncbi.nlm.nih.gov/entrez/eutils/esearch.fcgi"
    esummary_url = "https://eutils.ncbi.nlm.nih.gov/entrez/eutils/esummary.fcgi"
    efetch_url = "https://eutils.ncbi.nlm.nih.gov/entrez/eutils/efetch.fcgi"

    esearch_params = {
        "db": "pubmed", "term": query, "retmode": "json", "retmax": retmax,
        "sort": "relevance", "tool": "nutev", "email": email
    }
    data = safe_get(session, esearch_url, params=esearch_params).json()
    ids = data.get("esearchresult", {}).get("idlist", []) or []
    if not ids:
        return []

    esummary_params = {"db": "pubmed", "id": ",".join(ids), "retmode": "json", "tool": "nutev", "email": email}
    summary = safe_get(session, esummary_url, params=esummary_params).json()
    result_map = summary.get("result", {}) if isinstance(summary, dict) else {}

    efetch_params = {"db": "pubmed", "id": ",".join(ids), "retmode": "xml", "tool": "nutev", "email": email}
    xml_text = safe_get(session, efetch_url, params=efetch_params).text

    abstracts, dois, journals, languages, years, pmcids = {}, {}, {}, {}, {}, {}
    try:
        root = ET.fromstring(xml_text)
        for article in root.findall(".//PubmedArticle"):
            pmid = first_nonempty(article.findtext(".//PMID"))
            if not pmid:
                continue
            abs_parts = []
            for node in article.findall(".//Abstract/AbstractText"):
                label = node.attrib.get("Label", "")
                text = "".join(node.itertext()).strip()
                joined = f"{label}: {text}" if label else text
                if joined:
                    abs_parts.append(joined)
            abstracts[pmid] = " ".join(abs_parts)
            journals[pmid] = first_nonempty(article.findtext(".//Journal/Title"))
            languages[pmid] = first_nonempty(article.findtext(".//Language"))
            years[pmid] = first_nonempty(
                article.findtext(".//PubDate/Year"),
                article.findtext(".//ArticleDate/Year"),
                re.findall(r"(?:19|20)\d{2}", first_nonempty(article.findtext(".//PubDate/MedlineDate")))[0]
                if re.findall(r"(?:19|20)\d{2}", first_nonempty(article.findtext(".//PubDate/MedlineDate"))) else ""
            )
            for aid in article.findall(".//ArticleId"):
                idtype = (aid.attrib.get("IdType") or "").lower()
                if idtype == "doi":
                    dois[pmid] = (aid.text or "").strip()
                elif idtype == "pmc":
                    pmcids[pmid] = (aid.text or "").strip()
    except Exception as exc:
        log(f"PubMed parse warning: {exc}")

    rows = []
    for pmid in ids:
        doc = safe_dict(result_map.get(pmid))
        title = first_nonempty(doc.get("title"), "")
        year = first_nonempty(str(doc.get("pubdate", ""))[:4], years.get(pmid, ""))
        authors = "; ".join(a.get("name", "") for a in safe_list(doc.get("authors")) if safe_dict(a).get("name"))
        journal = first_nonempty(doc.get("fulljournalname"), journals.get(pmid, ""))
        doi = first_nonempty(dois.get(pmid), "")
        pmcid = first_nonempty(pmcids.get(pmid), "")
        landing = f"https://pubmed.ncbi.nlm.nih.gov/{pmid}/"
        rows.append(SearchRecord(
            uid=make_uid("pubmed", workstream, title, year, pmid),
            workstream=workstream,
            query_variant=query_variant,
            source="pubmed",
            title=title,
            abstract=abstracts.get(pmid, ""),
            authors=authors,
            year=year,
            language=languages.get(pmid, ""),
            doi=doi,
            pmid=pmid,
            pmcid=pmcid,
            journal=journal,
            publisher="NCBI / PubMed",
            document_type="journal article",
            landing_url=landing,
            query_used=query,
            collected_at_utc=now_utc_iso(),
        ))
    return rows

def search_europepmc(session: requests.Session, workstream: str, query_variant: str, query: str, retmax: int) -> List[SearchRecord]:
    url = "https://www.ebi.ac.uk/europepmc/webservices/rest/search"
    params = {"query": query, "format": "json", "pageSize": retmax, "sort": "RELEVANCE"}
    data = safe_get(session, url, params=params).json()
    results = data.get("resultList", {}).get("result", []) or []

    rows = []
    for item in results:
        title = first_nonempty(item.get("title"), "")
        year = first_nonempty(item.get("pubYear"), "")
        doi = first_nonempty(item.get("doi"), "")
        pmid = first_nonempty(item.get("pmid"), "")
        pmcid = first_nonempty(item.get("pmcid"), "")
        full_text_url = ""
        for ft in safe_list(safe_dict(item.get("fullTextUrlList")).get("fullTextUrl")):
            ft_url = safe_dict(ft).get("url", "")
            if ft_url:
                full_text_url = ft_url
                break

        rows.append(SearchRecord(
            uid=make_uid("europepmc", workstream, title, year, doi or pmid),
            workstream=workstream,
            query_variant=query_variant,
            source="europepmc",
            title=title,
            snippet=first_nonempty(item.get("authorString"), ""),
            abstract=first_nonempty(item.get("abstractText"), ""),
            authors=first_nonempty(item.get("authorString"), ""),
            year=year,
            language=first_nonempty(item.get("language"), ""),
            doi=doi,
            pmid=pmid,
            pmcid=pmcid,
            journal=first_nonempty(item.get("journalTitle"), ""),
            publisher="Europe PMC",
            document_type=first_nonempty(item.get("pubType"), ""),
            landing_url=first_nonempty(full_text_url, f"https://europepmc.org/article/MED/{pmid}" if pmid else ""),
            pdf_url=full_text_url if detect_extension_from_response(full_text_url, "") == ".pdf" else "",
            oa_url=full_text_url,
            is_oa=bool(full_text_url),
            query_used=query,
            collected_at_utc=now_utc_iso(),
        ))
    return rows

def search_openalex(session: requests.Session, workstream: str, query_variant: str, query: str, retmax: int) -> List[SearchRecord]:
    url = "https://api.openalex.org/works"
    params = {"search": query, "per-page": retmax, "mailto": "nutev@example.org"}
    data = safe_get(session, url, params=params).json()
    results = data.get("results", []) or []

    rows = []
    for item in results:
        title = first_nonempty(item.get("display_name"), "")
        year = first_nonempty(str(item.get("publication_year") or ""), "")
        authorships = safe_list(item.get("authorships"))
        authors = "; ".join(
            safe_dict(a.get("author")).get("display_name", "")
            for a in authorships if safe_dict(a).get("author")
        )
        primary_location = safe_dict(item.get("primary_location"))
        best_oa = safe_dict(item.get("best_oa_location"))
        open_access = safe_dict(item.get("open_access"))
        source_info = safe_dict(primary_location.get("source"))
        doi = (item.get("doi") or "").replace("https://doi.org/", "").strip()
        landing = first_nonempty(
            primary_location.get("landing_page_url"),
            best_oa.get("landing_page_url"),
            item.get("id")
        )
        pdf_url = first_nonempty(
            primary_location.get("pdf_url"),
            best_oa.get("pdf_url"),
            safe_dict(open_access).get("oa_url")
        )
        rows.append(SearchRecord(
            uid=make_uid("openalex", workstream, title, year, doi or landing),
            workstream=workstream,
            query_variant=query_variant,
            source="openalex",
            title=title,
            abstract=decode_openalex_abstract(item.get("abstract_inverted_index")),
            authors=authors,
            year=year,
            language=first_nonempty(item.get("language"), ""),
            doi=doi,
            journal=first_nonempty(source_info.get("display_name"), ""),
            publisher=first_nonempty(source_info.get("host_organization_name"), "OpenAlex"),
            document_type=first_nonempty(item.get("type"), ""),
            landing_url=landing,
            pdf_url=pdf_url if detect_extension_from_response(pdf_url, "application/pdf") == ".pdf" else "",
            oa_url=first_nonempty(pdf_url, landing),
            is_oa=bool(open_access.get("is_oa") or primary_location.get("is_oa") or best_oa.get("is_oa")),
            query_used=query,
            collected_at_utc=now_utc_iso(),
        ))
    return rows

def search_crossref(session: requests.Session, workstream: str, query_variant: str, query: str, email: str, retmax: int) -> List[SearchRecord]:
    url = "https://api.crossref.org/works"
    candidate_queries = []
    for q in [query, compact_free_text_query(query, 16), compact_free_text_query(query, 10)]:
        q = normalize_whitespace(q)
        if q and q not in candidate_queries:
            candidate_queries.append(q)

    last_exc = None
    data = None
    for idx, candidate in enumerate(candidate_queries):
        params = {"rows": retmax, "mailto": email}
        if idx == 0:
            params["query.bibliographic"] = candidate
        else:
            params["query"] = candidate
        try:
            data = safe_get(session, url, params=params).json()
            break
        except Exception as exc:
            last_exc = exc
            continue

    if data is None:
        raise last_exc if last_exc else RuntimeError("Crossref sem resposta")
    items = safe_list(safe_dict(data.get("message")).get("items"))

    rows = []
    for item in items:
        title = first_nonempty(*(item.get("title") or []))
        year = ""
        try:
            date_parts = safe_dict(item.get("issued")).get("date-parts", [[]])
            year = str(date_parts[0][0]) if date_parts and date_parts[0] else ""
        except Exception:
            year = ""
        authors = "; ".join(
            " ".join(filter(None, [safe_dict(a).get("given", ""), safe_dict(a).get("family", "")])).strip()
            for a in safe_list(item.get("author")) if a
        )
        doi = first_nonempty(item.get("DOI"), "")
        abstract = re.sub(r"<[^>]+>", " ", first_nonempty(item.get("abstract"), "")) if item.get("abstract") else ""
        landing = first_nonempty(item.get("URL"), "")
        rows.append(SearchRecord(
            uid=make_uid("crossref", workstream, title, year, doi or landing),
            workstream=workstream,
            query_variant=query_variant,
            source="crossref",
            title=title,
            abstract=abstract,
            authors=authors,
            year=year,
            language=first_nonempty(item.get("language"), ""),
            doi=doi,
            journal=first_nonempty(*(item.get("container-title") or [""])),
            publisher=first_nonempty(item.get("publisher"), "Crossref"),
            document_type=first_nonempty(item.get("type"), ""),
            landing_url=landing,
            query_used=query,
            collected_at_utc=now_utc_iso(),
        ))
    return rows

def search_doaj(session: requests.Session, workstream: str, query_variant: str, query: str, retmax: int) -> List[SearchRecord]:
    safe_q = quote(query)
    url = f"https://doaj.org/api/search/articles/{safe_q}"
    params = {"pageSize": retmax}
    data = safe_get(session, url, params=params).json()
    results = data.get("results", []) or []

    rows = []
    for item in results:
        bib = safe_dict(item.get("bibjson"))
        title = first_nonempty(bib.get("title"), "")
        year = first_nonempty(str(bib.get("year") or ""), "")
        authors = "; ".join(safe_dict(a).get("name", "") for a in safe_list(bib.get("author")) if safe_dict(a).get("name"))
        link_urls = [safe_dict(l).get("url") for l in safe_list(bib.get("link")) if safe_dict(l).get("url")]
        fulltext = next((safe_dict(l).get("url") for l in safe_list(bib.get("link")) if (safe_dict(l).get("type") or "").lower() == "fulltext"), "")
        doi = ""
        for ident in safe_list(bib.get("identifier")):
            if (safe_dict(ident).get("type") or "").lower() == "doi":
                doi = safe_dict(ident).get("id", "")
                break
        rows.append(SearchRecord(
            uid=make_uid("doaj", workstream, title, year, doi or "".join(link_urls[:1])),
            workstream=workstream,
            query_variant=query_variant,
            source="doaj",
            title=title,
            abstract=first_nonempty(bib.get("abstract"), ""),
            authors=authors,
            year=year,
            language=first_nonempty(safe_list(bib.get("language"))[0] if safe_list(bib.get("language")) else "", ""),
            doi=doi,
            journal=first_nonempty(safe_dict(bib.get("journal")).get("title"), ""),
            publisher=first_nonempty(safe_dict(bib.get("journal")).get("publisher"), "DOAJ"),
            document_type=first_nonempty(bib.get("type"), "article"),
            landing_url=first_nonempty(*link_urls),
            pdf_url=fulltext if detect_extension_from_response(fulltext, "") == ".pdf" else "",
            oa_url=fulltext,
            is_oa=bool(fulltext or link_urls),
            query_used=query,
            collected_at_utc=now_utc_iso(),
        ))
    return rows

def enrich_with_unpaywall(session: requests.Session, records: List[SearchRecord], email: str, max_records: int = 250) -> List[SearchRecord]:
    count = 0
    for record in records:
        if count >= max_records:
            break
        if not record.doi or record.is_oa:
            continue
        doi = quote(record.doi, safe="")
        url = f"https://api.unpaywall.org/v2/{doi}"
        try:
            data = safe_get(session, url, params={"email": email}).json()
        except Exception:
            continue
        record.is_oa = bool(data.get("is_oa"))
        best = safe_dict(data.get("best_oa_location"))
        record.oa_url = first_nonempty(best.get("url_for_pdf"), best.get("url_for_landing_page"), record.oa_url)
        record.pdf_url = first_nonempty(best.get("url_for_pdf"), record.pdf_url)
        count += 1
    return records


# ============================================================
# SEARCH WEB - GOOGLE + FALLBACK HTML
# ============================================================

def search_google_cse(session: requests.Session, workstream: str, query_variant: str, query: str, api_key: str, cse_id: str, retmax: int) -> List[SearchRecord]:
    url = "https://customsearch.googleapis.com/customsearch/v1"
    num = min(max(1, retmax), 10)
    params = {"key": api_key, "cx": cse_id, "q": query, "num": num}
    data = safe_get(session, url, params=params).json()
    items = safe_list(data.get("items"))

    rows = []
    for item in items:
        title = decode_html_entities(first_nonempty(item.get("title"), ""))
        snippet = decode_html_entities(first_nonempty(item.get("snippet"), ""))
        link = first_nonempty(item.get("link"), "")
        file_hint = detect_extension_from_response(link, "")
        rows.append(SearchRecord(
            uid=make_uid("google_cse", workstream, title, "", link),
            workstream=workstream,
            query_variant=query_variant,
            source="google_cse",
            record_kind="web",
            title=title,
            snippet=snippet,
            document_type="web_result",
            landing_url=link,
            pdf_url=link if file_hint == ".pdf" else "",
            file_ext_hint=file_hint,
            is_oa=file_hint == ".pdf",
            query_used=query,
            collected_at_utc=now_utc_iso(),
        ))
    return rows

def search_ddg_html(session: requests.Session, workstream: str, query_variant: str, query: str, retmax: int) -> List[SearchRecord]:
    if BeautifulSoup is None:
        return []
    url = "https://html.duckduckgo.com/html/"
    resp = safe_post(session, url, data={"q": query})
    soup = BeautifulSoup(resp.text, "html.parser")

    rows = []
    for a in soup.select("a.result__a, a[data-testid='result-title-a']"):
        href = normalize_url_from_search(a.get("href", ""))
        title = decode_html_entities(a.get_text(" ", strip=True))
        if not href or not title:
            continue
        container = a.parent
        snippet = ""
        if container:
            snip = container.find_next(string=False)
        # fallback simple search for nearest snippet element
        parent = a.parent
        if parent and hasattr(parent, "find"):
            snip_tag = parent.find(class_=re.compile("result__snippet"))
            if snip_tag:
                snippet = snip_tag.get_text(" ", strip=True)
        file_hint = detect_extension_from_response(href, "")
        rows.append(SearchRecord(
            uid=make_uid("ddg_web", workstream, title, "", href),
            workstream=workstream,
            query_variant=query_variant,
            source="ddg_web",
            record_kind="web",
            title=title,
            snippet=snippet,
            document_type="web_result",
            landing_url=href,
            pdf_url=href if file_hint == ".pdf" else "",
            file_ext_hint=file_hint,
            is_oa=file_hint == ".pdf",
            query_used=query,
            collected_at_utc=now_utc_iso(),
        ))
        if len(rows) >= retmax:
            break
    return rows


# ============================================================
# OFFICIAL CRAWLER E CANDIDATOS
# ============================================================

def crawl_official_sources(session: requests.Session, manifest: dict, project_root: Path, max_depth: int = 1, max_links_per_seed: int = 80) -> pd.DataFrame:
    rows = []
    if BeautifulSoup is None:
        log("BeautifulSoup não disponível; pulando crawler de sites oficiais.")
        return pd.DataFrame()

    for source in manifest.get("sources", []):
        source_name = source.get("name", "official_source")
        allowed_domains = source.get("allowed_domains", [])
        seed_urls = source.get("seed_urls", [])
        workstream = source.get("workstream", "")
        for seed in seed_urls:
            queue = [(seed, 0)]
            seen = set()
            found = 0
            while queue and found < max_links_per_seed:
                current_url, depth = queue.pop(0)
                if current_url in seen or depth > max_depth:
                    continue
                seen.add(current_url)
                try:
                    resp = safe_get(session, current_url)
                except Exception:
                    continue
                ctype = (resp.headers.get("Content-Type") or "").lower()
                if "text/html" not in ctype:
                    continue

                soup = BeautifulSoup(resp.text, "html.parser")
                page_title = soup.title.get_text(" ", strip=True) if soup.title else ""
                text_blob = soup.get_text(" ", strip=True)[:1000]

                for a in soup.find_all("a", href=True):
                    href = urljoin(current_url, a["href"])
                    anchor_text = a.get_text(" ", strip=True)
                    if not is_allowed_domain(href, allowed_domains):
                        continue
                    if not looks_like_relevant_link(href, anchor_text):
                        continue
                    ext = detect_extension_from_response(href, "")
                    source_bucket = "document_link" if ext in DOWNLOADABLE_EXT else "landing_page"
                    rows.append({
                        "workstream": workstream,
                        "source_name": source_name,
                        "seed_url": seed,
                        "page_url": current_url,
                        "page_title": page_title,
                        "page_text_excerpt": text_blob[:400],
                        "link_text": anchor_text,
                        "link_url": href,
                        "link_ext": ext,
                        "link_bucket": source_bucket,
                        "collected_at_utc": now_utc_iso(),
                    })
                    found += 1
                    if source_bucket == "landing_page" and depth < max_depth and found < max_links_per_seed:
                        queue.append((href, depth + 1))

    df = pd.DataFrame(rows)
    if not df.empty:
        path = build_default_paths(project_root)["candidate_links_csv"]
        ensure_folder(path.parent)
        mode = "a" if path.exists() else "w"
        header = not path.exists()
        df.to_csv(path, mode=mode, header=header, index=False, encoding="utf-8")
    return df


# ============================================================
# DOWNLOADERS E WEB CRAWL
# ============================================================

def append_download_manifest(project_root: Path, rows: List[DownloadManifestRow]) -> None:
    if not rows:
        return
    path = build_default_paths(project_root)["download_manifest_csv"]
    ensure_folder(path.parent)
    df = as_dataframe(rows, DownloadManifestRow)
    mode = "a" if path.exists() else "w"
    header = not path.exists()
    df.to_csv(path, mode=mode, header=header, index=False, encoding="utf-8")

def download_single_public_url(session: requests.Session, record: SearchRecord, bucket_root: Path, bucket_name: str) -> DownloadManifestRow:
    url = first_nonempty(record.pdf_url, record.oa_url, record.landing_url)
    ext_guess = detect_extension_from_response(url, "")
    if ext_guess not in DOWNLOADABLE_EXT:
        return DownloadManifestRow(record.workstream, bucket_name, record.source, record.title, url, record.parent_url, "", "skipped", "URL não parece arquivo público baixável", 0, "", now_utc_iso())
    out_path = bucket_root / filename_from_record(record, ext_guess, bucket_name)
    if out_path.exists() and out_path.stat().st_size > 0:
        return DownloadManifestRow(record.workstream, bucket_name, record.source, record.title, url, record.parent_url, str(out_path), "exists", "", out_path.stat().st_size, "", now_utc_iso())

    ok, info, size = stream_download(session, url, out_path)
    status = "downloaded" if ok else "failed"
    if not ok and out_path.exists():
        try:
            out_path.unlink()
        except Exception:
            pass
    return DownloadManifestRow(record.workstream, bucket_name, record.source, record.title, url, record.parent_url, str(out_path), status, info, size, info if ok else "", now_utc_iso())

def download_oa_records(session: requests.Session, records: List[SearchRecord], project_root: Path, max_downloads: int = DEFAULT_DOWNLOAD_MAX) -> pd.DataFrame:
    rows: List[DownloadManifestRow] = []
    downloaded = 0
    seen = set()
    for record in records:
        if downloaded >= max_downloads:
            break
        target_url = first_nonempty(record.pdf_url, record.oa_url)
        if not target_url or target_url in seen:
            continue
        ext_guess = detect_extension_from_response(target_url, "")
        if ext_guess != ".pdf":
            continue
        seen.add(target_url)
        ws_dir = ensure_folder(project_root / "03_corpus" / "03A_bibliographic_oa" / record.workstream)
        row = download_single_public_url(session, record, ws_dir, "bibliographic_oa")
        rows.append(row)
        if row.status in {"downloaded", "exists"}:
            downloaded += 1
    append_download_manifest(project_root, rows)
    return as_dataframe(rows, DownloadManifestRow)

def save_html_snapshot(project_root: Path, workstream: str, source: str, url: str, html: str, title: str = "") -> Path:
    ws_dir = ensure_folder(project_root / "03_corpus" / "03C_web_landing_pages" / workstream)
    out_path = ws_dir / url_to_html_filename(workstream, source, url, title=title)
    if not out_path.exists():
        out_path.write_text(html, encoding="utf-8", errors="ignore")
    return out_path

def crawl_result_pages_and_collect_candidates(
    session: requests.Session,
    search_records: List[SearchRecord],
    project_root: Path,
    max_pages: int = DEFAULT_CRAWL_PAGES,
    max_links_per_page: int = DEFAULT_CRAWL_LINKS_PER_PAGE,
) -> pd.DataFrame:
    if BeautifulSoup is None:
        return pd.DataFrame()
    rows = []
    seen_pages = set()
    pages_done = 0
    for record in search_records:
        if pages_done >= max_pages:
            break
        url = first_nonempty(record.landing_url, record.oa_url)
        if not url or url in seen_pages:
            continue
        seen_pages.add(url)

        # se a própria landing já for arquivo
        ext_guess = detect_extension_from_response(url, "")
        if ext_guess in DOWNLOADABLE_EXT:
            rows.append({
                "workstream": record.workstream,
                "source_name": record.source,
                "seed_url": record.query_used,
                "page_url": record.parent_url,
                "page_title": record.title,
                "page_text_excerpt": record.snippet[:400],
                "link_text": record.title,
                "link_url": url,
                "link_ext": ext_guess,
                "link_bucket": "document_link",
                "collected_at_utc": now_utc_iso(),
            })
            pages_done += 1
            continue

        try:
            resp = safe_get(session, url)
        except Exception:
            continue
        ctype = (resp.headers.get("Content-Type") or "").lower()
        if "text/html" not in ctype:
            continue
        html = resp.text
        soup = BeautifulSoup(html, "html.parser")
        page_title = soup.title.get_text(" ", strip=True) if soup.title else record.title
        page_text_excerpt = soup.get_text(" ", strip=True)[:600]
        save_html_snapshot(project_root, record.workstream, record.source, url, html, title=page_title)

        local_count = 0
        for a in soup.find_all("a", href=True):
            href = urljoin(url, a["href"])
            href_low = href.lower()
            if any(h in href_low for h in PAYWALL_HINTS):
                continue
            anchor_text = a.get_text(" ", strip=True)
            if not looks_like_relevant_link(href, anchor_text):
                continue
            ext = detect_extension_from_response(href, "")
            if ext not in DOWNLOADABLE_EXT and ext != ".html":
                continue
            rows.append({
                "workstream": record.workstream,
                "source_name": record.source,
                "seed_url": record.query_used,
                "page_url": url,
                "page_title": page_title,
                "page_text_excerpt": page_text_excerpt[:400],
                "link_text": anchor_text,
                "link_url": href,
                "link_ext": ext,
                "link_bucket": "document_link" if ext in DOWNLOADABLE_EXT else "landing_page",
                "collected_at_utc": now_utc_iso(),
            })
            local_count += 1
            if local_count >= max_links_per_page:
                break

        pages_done += 1

    df = pd.DataFrame(rows)
    if not df.empty:
        path = build_default_paths(project_root)["candidate_links_csv"]
        ensure_folder(path.parent)
        mode = "a" if path.exists() else "w"
        header = not path.exists()
        df.to_csv(path, mode=mode, header=header, index=False, encoding="utf-8")
    return df

def download_candidate_links(
    session: requests.Session,
    candidate_df: pd.DataFrame,
    project_root: Path,
    bucket_name: str,
    max_downloads: int = DEFAULT_DOWNLOAD_MAX,
) -> pd.DataFrame:
    rows: List[DownloadManifestRow] = []
    downloaded = 0
    seen = set()
    if candidate_df is None or candidate_df.empty:
        return pd.DataFrame()
    for _, row in candidate_df.iterrows():
        if downloaded >= max_downloads:
            break
        url = first_nonempty(row.get("link_url"), "")
        if not url or url in seen:
            continue
        seen.add(url)
        ext = detect_extension_from_response(url, first_nonempty(row.get("link_ext"), ""))
        if ext not in DOWNLOADABLE_EXT:
            continue
        workstream = first_nonempty(row.get("workstream"), "misc")
        source_name = slugify(first_nonempty(row.get("source_name"), "web"), 18)
        bucket_folder = "03B_web_direct_docs" if bucket_name == "web_direct_docs" else "03D_official_seed_docs"
        ws_dir = ensure_folder(project_root / "03_corpus" / bucket_folder / workstream)
        rec = SearchRecord(
            uid=stable_hash(url)[:16],
            workstream=workstream,
            query_variant="download",
            source=source_name,
            record_kind="web",
            title=first_nonempty(row.get("link_text"), row.get("page_title"), "documento"),
            snippet=first_nonempty(row.get("page_text_excerpt"), ""),
            landing_url=url,
            file_ext_hint=ext,
            parent_url=first_nonempty(row.get("page_url"), ""),
            query_used=first_nonempty(row.get("seed_url"), ""),
            collected_at_utc=now_utc_iso(),
        )
        dl_row = download_single_public_url(session, rec, ws_dir, bucket_name)
        rows.append(dl_row)
        if dl_row.status in {"downloaded", "exists"}:
            downloaded += 1
    append_download_manifest(project_root, rows)
    return as_dataframe(rows, DownloadManifestRow)


# ============================================================
# LOCAL DOCUMENTS: OCR / EXTRAÇÃO
# ============================================================

def choose_source_bucket(file_path: Path) -> str:
    low = str(file_path).lower()
    if "03a_bibliographic_oa" in low:
        return "bibliographic_oa"
    if "03b_web_direct_docs" in low:
        return "web_direct_doc"
    if "03c_web_landing_pages" in low:
        return "web_landing_html"
    if "03d_official_seed_docs" in low:
        return "official_seed_doc"
    if "03e_manual_drop" in low:
        return "manual_drop"
    return "unspecified"

def infer_workstream_from_path(file_path: Path) -> str:
    low = str(file_path).lower()
    for ws in ["busca1", "busca2a", "busca2b", "artigo3_framework"]:
        if ws in low:
            return ws
    return "artigo3_framework"

def extract_text_pdf(file_path: Path, ocr_force: bool = False, max_ocr_pages: int = 60) -> Tuple[str, bool]:
    if pdfplumber is None:
        return "ERROR: pdfplumber não instalado", False
    full_text, used_ocr = [], False
    try:
        with pdfplumber.open(str(file_path)) as pdf:
            for i, page in enumerate(pdf.pages):
                text = page.extract_text() or ""
                needs_ocr = ocr_force or len(text.strip()) < 30
                if needs_ocr and pytesseract is not None and Image is not None and i < max_ocr_pages:
                    try:
                        img = page.to_image(resolution=OCR_RESOLUTION).original
                        ocr_text = pytesseract.image_to_string(img, lang="por+eng+spa")
                        if len(ocr_text.strip()) > len(text.strip()):
                            text = ocr_text
                            used_ocr = True
                    except Exception:
                        pass
                if text:
                    full_text.append(text)
    except Exception as exc:
        return f"ERROR_PDF: {exc}", False
    return "\n".join(full_text), used_ocr

def extract_text_image(file_path: Path) -> Tuple[str, bool]:
    if pytesseract is None or Image is None:
        return "ERROR: pytesseract/Pillow não instalados", False
    try:
        img = Image.open(file_path)
        txt = pytesseract.image_to_string(img, lang="por+eng+spa")
        return txt, True
    except Exception as exc:
        return f"ERROR_IMAGE: {exc}", False

def extract_text_docx(file_path: Path) -> Tuple[str, bool]:
    if py_docx is None:
        return "ERROR: python-docx não instalado", False
    try:
        doc = py_docx.Document(str(file_path))
        chunks = []
        for p in doc.paragraphs:
            if p.text.strip():
                chunks.append(p.text.strip())
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    chunks.append(row_text)
        return "\n".join(chunks), False
    except Exception as exc:
        return f"ERROR_DOCX: {exc}", False

def extract_text_pptx(file_path: Path) -> Tuple[str, bool]:
    try:
        chunks = []
        with zipfile.ZipFile(file_path, "r") as zf:
            for name in zf.namelist():
                if name.startswith("ppt/slides/slide") and name.endswith(".xml"):
                    xml = zf.read(name).decode("utf-8", errors="ignore")
                    texts = re.findall(r"<a:t>(.*?)</a:t>", xml)
                    if texts:
                        chunks.append(" ".join(decode_html_entities(t) for t in texts))
        return "\n".join(chunks), False
    except Exception as exc:
        return f"ERROR_PPTX: {exc}", False

def extract_text_spreadsheet(file_path: Path) -> Tuple[str, bool]:
    try:
        if file_path.suffix.lower() == ".csv":
            df = pd.read_csv(file_path)
        else:
            df = pd.read_excel(file_path)
        return df.astype(str).to_string(index=False), False
    except Exception as exc:
        return f"ERROR_SHEET: {exc}", False

def extract_text_html(file_path: Path) -> Tuple[str, bool]:
    try:
        html = file_path.read_text(encoding="utf-8", errors="ignore")
        return html_to_text(html), False
    except Exception as exc:
        return f"ERROR_HTML: {exc}", False

def extract_text_local(file_path: Path, ocr_force: bool = False) -> Tuple[str, bool]:
    ext = file_path.suffix.lower()
    if ext == ".pdf":
        return extract_text_pdf(file_path, ocr_force=ocr_force)
    if ext == ".docx":
        return extract_text_docx(file_path)
    if ext == ".pptx":
        return extract_text_pptx(file_path)
    if ext in [".csv", ".xlsx", ".xls"]:
        return extract_text_spreadsheet(file_path)
    if ext in [".html", ".htm"]:
        return extract_text_html(file_path)
    if ext in [".png", ".jpg", ".jpeg", ".tif", ".tiff", ".webp"]:
        return extract_text_image(file_path)
    if ext == ".json":
        try:
            payload = json.loads(file_path.read_text(encoding="utf-8", errors="ignore"))
            return json.dumps(payload, ensure_ascii=False, indent=2), False
        except Exception:
            return file_path.read_text(encoding="utf-8", errors="ignore"), False
    try:
        return file_path.read_text(encoding="utf-8", errors="ignore"), False
    except Exception as exc:
        return f"ERROR_TEXT: {exc}", False

def extract_excerpt_for_terms(text: str, terms: Sequence[str], window: int = 180, max_hits: int = 2) -> str:
    low = text.lower()
    excerpts = []
    for term in terms:
        idx = low.find(term.lower())
        if idx >= 0:
            start = max(0, idx - window)
            end = min(len(text), idx + len(term) + window)
            excerpts.append(text[start:end].replace("\n", " ").strip())
        if len(excerpts) >= max_hits:
            break
    return " || ".join(excerpts) if excerpts else ""

def detect_local_domains(text: str, taxonomy: dict) -> Dict[str, str]:
    mappings = {
        "LM_Pillars": sum(taxonomy["global"]["lifestyle_medicine_pillars"].values(), []),
        "Diet_Patterns": sum(taxonomy["global"]["diet_patterns"].values(), []),
        "Nutrition_Composition": sum(taxonomy["global"]["nutrition_domains"].values(), []),
        "Implementation_Behavior": sum(taxonomy["global"]["implementation_behavior"].values(), []),
        "Clinical_Conditions": sum(taxonomy["clinical"].values(), []),
        "Outcomes": sum(taxonomy["outcomes"].values(), []),
    }
    out = {}
    for key, terms in mappings.items():
        excerpt = extract_excerpt_for_terms(text, terms)
        out[key] = excerpt or "NÃO REPORTADO"
    return out

def analyze_local_documents(project_root: Path, taxonomy: dict, ocr_force: bool = False) -> Tuple[pd.DataFrame, Dict[str, pd.DataFrame]]:
    corpus_root = project_root / "03_corpus"
    files = [p for p in corpus_root.rglob("*") if p.is_file() and p.suffix.lower() in SUPPORTED_LOCAL_EXT]
    rows: List[LocalDocumentRecord] = []
    domain_rows = []

    for fp in files:
        try:
            st = fp.stat()
            text, used_ocr = extract_text_local(fp, ocr_force=ocr_force)
            text = normalize_whitespace(text)
            inferred_workstream = infer_workstream_from_path(fp)
            source_bucket = choose_source_bucket(fp)
            doc_id = stable_hash(str(fp), str(st.st_mtime), str(st.st_size))[:18]
            bank = collect_term_bank(taxonomy, inferred_workstream)
            local_text = " ".join([fp.name, text[:12000]])
            matched_domains = term_hits(local_text, bank.get("nutrition_terms", []) + bank.get("implementation_terms", []))
            matched_conditions = term_hits(local_text, bank.get("clinical_terms", []) + bank.get("condition_terms", []))
            matched_patterns = term_hits(local_text, bank.get("pattern_terms", []))
            matched_outcomes = term_hits(local_text, bank.get("outcome_terms", []))

            score = 0.0
            score += min(len(matched_domains), 6) * 1.3
            score += min(len(matched_conditions), 5) * 1.7
            score += min(len(matched_patterns), 5) * 1.5
            score += min(len(matched_outcomes), 5) * 1.2

            if score >= 8.0:
                prisma_decision = "INCLUDE"
                prisma_reason = "Documento local com alta aderência terminológica"
            elif score >= 3.8:
                prisma_decision = "UNCERTAIN"
                prisma_reason = "Documento local parcialmente aderente; revisar"
            else:
                prisma_decision = "EXCLUDE"
                prisma_reason = "Baixa aderência terminológica"

            domain_evidence = detect_local_domains(text, taxonomy)
            domain_rows.append({
                "doc_id": doc_id,
                "file_name": fp.name,
                "inferred_workstream": inferred_workstream,
                **domain_evidence,
            })

            rows.append(LocalDocumentRecord(
                doc_id=doc_id,
                file_path=str(fp),
                file_name=fp.name,
                file_ext=fp.suffix.lower(),
                file_size_kb=round(st.st_size / 1024, 2),
                modified_time=datetime.fromtimestamp(st.st_mtime).isoformat(timespec="seconds"),
                inferred_workstream=inferred_workstream,
                source_bucket=source_bucket,
                prisma_decision=prisma_decision,
                prisma_reason=prisma_reason,
                language_hint=infer_language_hint(text),
                used_ocr=used_ocr,
                text_len=len(text),
                text_hash=stable_hash(text) if text else "",
                dedup_group="",
                duplicate_of="",
                title_guess=fp.stem,
                matched_domains="; ".join(matched_domains[:20]),
                matched_conditions="; ".join(matched_conditions[:20]),
                matched_patterns="; ".join(matched_patterns[:20]),
                matched_outcomes="; ".join(matched_outcomes[:20]),
                error="" if not text.startswith("ERROR") else text[:200],
                raw_excerpt=text[:3000],
            ))
        except Exception as exc:
            rows.append(LocalDocumentRecord(
                doc_id=stable_hash(str(fp))[:18],
                file_path=str(fp),
                file_name=fp.name,
                file_ext=fp.suffix.lower(),
                file_size_kb=0,
                modified_time="",
                inferred_workstream=infer_workstream_from_path(fp),
                source_bucket=choose_source_bucket(fp),
                prisma_decision="ERROR",
                prisma_reason="Falha de processamento",
                language_hint="",
                used_ocr=False,
                text_len=0,
                text_hash="",
                dedup_group="",
                duplicate_of="",
                title_guess=fp.stem,
                matched_domains="",
                matched_conditions="",
                matched_patterns="",
                matched_outcomes="",
                error=str(exc),
                raw_excerpt="",
            ))

    df = as_dataframe(rows, LocalDocumentRecord)

    if TfidfVectorizer is not None and cosine_similarity is not None and len(df) > 1:
        candidate_idx = df.index[df["text_len"].fillna(0).astype(int) >= LOCAL_MIN_TEXT_FOR_DEDUP].tolist()
        if len(candidate_idx) >= 2:
            texts = df.loc[candidate_idx, "raw_excerpt"].fillna("").astype(str).tolist()
            X = TfidfVectorizer(min_df=1, ngram_range=(1, 2)).fit_transform(texts)
            sim = cosine_similarity(X)
            group_counter = 0
            parent = list(range(len(candidate_idx)))

            def find(a):
                while parent[a] != a:
                    parent[a] = parent[parent[a]]
                    a = parent[a]
                return a

            def union(a, b):
                ra, rb = find(a), find(b)
                if ra != rb:
                    parent[rb] = ra

            for i in range(len(candidate_idx)):
                for j in range(i + 1, len(candidate_idx)):
                    if sim[i, j] >= 0.93:
                        union(i, j)

            group_map = {}
            for pos, idx in enumerate(candidate_idx):
                root = find(pos)
                if root not in group_map:
                    group_counter += 1
                    group_map[root] = f"G{group_counter:04d}"
                df.loc[idx, "dedup_group"] = group_map[root]

            for _, gdf in df[df["dedup_group"] != ""].groupby("dedup_group"):
                rep_idx = gdf["text_len"].astype(int).idxmax()
                rep_doc = df.loc[rep_idx, "doc_id"]
                for idx in gdf.index:
                    if idx != rep_idx:
                        df.loc[idx, "duplicate_of"] = rep_doc

    summary_workstream = df.groupby(["inferred_workstream", "prisma_decision"], dropna=False).size().reset_index(name="n")
    summary_bucket = df.groupby(["source_bucket", "prisma_decision"], dropna=False).size().reset_index(name="n")
    domain_df = pd.DataFrame(domain_rows)

    domain_counts = []
    if not domain_df.empty:
        for col in [c for c in domain_df.columns if c not in {"doc_id", "file_name", "inferred_workstream"}]:
            n = int((domain_df[col] != "NÃO REPORTADO").sum())
            pct = float((domain_df[col] != "NÃO REPORTADO").mean())
            domain_counts.append({"domain": col, "n_documents": n, "pct_documents": pct})
    domain_counts_df = pd.DataFrame(domain_counts).sort_values("n_documents", ascending=False) if domain_counts else pd.DataFrame()

    return df, {
        "SUMMARY_WORKSTREAM": summary_workstream,
        "SUMMARY_BUCKET": summary_bucket,
        "DOMAIN_EVIDENCE": domain_df,
        "DOMAIN_COUNTS": domain_counts_df,
    }


# ============================================================
# EXPORTS
# ============================================================

def records_to_export_dfs(records: List[SearchRecord]) -> Dict[str, pd.DataFrame]:
    df = as_dataframe(records, SearchRecord)
    if df.empty:
        return {
            "MASTER": df,
            "SUMMARY_SOURCE": pd.DataFrame(),
            "SUMMARY_WORKSTREAM": pd.DataFrame(),
            "SUMMARY_KIND": pd.DataFrame(),
            "TOP_RECORDS": pd.DataFrame(),
        }
    summary_source = df.groupby(["source", "prisma_decision"], dropna=False).size().reset_index(name="n")
    summary_workstream = df.groupby(["workstream", "prisma_decision"], dropna=False).size().reset_index(name="n")
    summary_kind = df.groupby(["record_kind", "source"], dropna=False).size().reset_index(name="n")
    top_records = df.sort_values(["relevance_score", "year"], ascending=[False, False]).head(300)
    return {
        "MASTER": df,
        "SUMMARY_SOURCE": summary_source,
        "SUMMARY_WORKSTREAM": summary_workstream,
        "SUMMARY_KIND": summary_kind,
        "TOP_RECORDS": top_records,
    }

def build_rayyan_ready(records: List[SearchRecord]) -> pd.DataFrame:
    rows = []
    for r in records:
        rows.append({
            "id": r.uid,
            "title": r.title,
            "abstract": r.abstract or r.snippet,
            "authors": r.authors,
            "year": r.year,
            "journal": r.journal,
            "doi": r.doi,
            "source": r.source,
            "workstream": r.workstream,
            "document_type": r.document_type,
            "keywords_matched": r.keywords_matched,
            "relevance_score": r.relevance_score,
            "prisma_decision_auto": r.prisma_decision,
            "prisma_reason_auto": r.prisma_reason,
            "url": r.landing_url or r.oa_url,
        })
    return pd.DataFrame(rows)

def build_naming_example(record: SearchRecord) -> str:
    ext = ".pdf" if record.pdf_url else (record.file_ext_hint or ".html")
    return filename_from_record(record, ext, bucket="example")


# ============================================================
# MAIN SEARCH RUNNER
# ============================================================

def run_search_layers(
    project_root: Path,
    taxonomy: dict,
    query_pack: dict,
    email: str,
    retmax: int,
    web_retmax: int,
    enable_web: bool,
    enable_google: bool,
    enable_official_crawl: bool,
    google_api_key: str,
    google_cse_id: str,
    seed_manifest: dict,
) -> Tuple[List[SearchRecord], pd.DataFrame]:
    session = session_factory(email=email)
    all_records: List[SearchRecord] = []

    for workstream_key, payload in query_pack["workstreams"].items():
        source_map = payload["queries"]
        priority_sources = taxonomy["workstreams"][workstream_key].get("source_priority", [])
        ordered_sources = [s for s in ["pubmed", "europepmc", "openalex", "crossref", "doaj"] if s in source_map]
        ordered_sources = sorted(ordered_sources, key=lambda s: priority_sources.index(s) if s in priority_sources else 999)

        for source in ordered_sources:
            for query_variant, query in source_map[source].items():
                if not query.strip():
                    continue
                log(f"Buscando {workstream_key} | {source} | {query_variant}")
                try:
                    if source == "pubmed":
                        recs = search_pubmed(session, workstream_key, query_variant, query, email, retmax)
                    elif source == "europepmc":
                        recs = search_europepmc(session, workstream_key, query_variant, query, retmax)
                    elif source == "openalex":
                        recs = search_openalex(session, workstream_key, query_variant, query, retmax)
                    elif source == "crossref":
                        recs = search_crossref(session, workstream_key, query_variant, query, email, retmax)
                    elif source == "doaj":
                        recs = search_doaj(session, workstream_key, query_variant, query, retmax)
                    else:
                        recs = []

                    for r in recs:
                        score, matched, decision, reason = score_search_record(r, taxonomy)
                        r.relevance_score = score
                        r.keywords_matched = "; ".join(matched)
                        r.prisma_decision = decision
                        r.prisma_reason = reason
                        r.source_rank = priority_sources.index(r.source) + 1 if r.source in priority_sources else 999
                    all_records.extend(recs)
                except Exception as exc:
                    log(f"Falha em {source} ({workstream_key}/{query_variant}): {exc}")

        if enable_web and "web" in source_map:
            for query_variant, query in source_map["web"].items():
                if not query.strip():
                    continue
                if enable_google and google_api_key and google_cse_id:
                    try:
                        log(f"Buscando {workstream_key} | google_cse | {query_variant}")
                        recs = search_google_cse(session, workstream_key, query_variant, query, google_api_key, google_cse_id, web_retmax)
                        for r in recs:
                            score, matched, decision, reason = score_search_record(r, taxonomy)
                            r.relevance_score = score
                            r.keywords_matched = "; ".join(matched)
                            r.prisma_decision = decision
                            r.prisma_reason = reason
                            r.source_rank = priority_sources.index(r.source) + 1 if r.source in priority_sources else 999
                        all_records.extend(recs)
                    except Exception as exc:
                        log(f"Falha em google_cse ({workstream_key}/{query_variant}): {exc}")
                try:
                    log(f"Buscando {workstream_key} | ddg_web | {query_variant}")
                    recs = search_ddg_html(session, workstream_key, query_variant, query, web_retmax)
                    for r in recs:
                        score, matched, decision, reason = score_search_record(r, taxonomy)
                        r.relevance_score = score
                        r.keywords_matched = "; ".join(matched)
                        r.prisma_decision = decision
                        r.prisma_reason = reason
                        r.source_rank = priority_sources.index(r.source) + 1 if r.source in priority_sources else 999
                    all_records.extend(recs)
                except Exception as exc:
                    log(f"Falha em ddg_web ({workstream_key}/{query_variant}): {exc}")

    try:
        all_records = enrich_with_unpaywall(session, all_records, email=email, max_records=250)
    except Exception as exc:
        log(f"Enriquecimento Unpaywall falhou: {exc}")

    official_df = pd.DataFrame()
    if enable_official_crawl:
        try:
            official_df = crawl_official_sources(session, seed_manifest, project_root, max_depth=1, max_links_per_seed=80)
            log(f"Crawler oficial encontrou {len(official_df)} links candidatos.")
        except Exception as exc:
            log(f"Crawler oficial falhou: {exc}")

    return deduplicate_records(all_records), official_df


# ============================================================
# WRITE OUTPUTS
# ============================================================

def write_outputs(
    project_root: Path,
    metadata_records: List[SearchRecord],
    local_df: pd.DataFrame | None,
    local_sheets: Dict[str, pd.DataFrame] | None,
) -> Dict[str, str]:
    paths = build_default_paths(project_root)
    output_registry: Dict[str, str] = {}

    if metadata_records:
        export_dfs = records_to_export_dfs(metadata_records)
        row_dicts_to_excel(export_dfs, paths["search_master_xlsx"])
        export_dfs["MASTER"].to_csv(paths["search_master_csv"], index=False, encoding="utf-8")
        rayyan_df = build_rayyan_ready(metadata_records)
        rayyan_df.to_csv(paths["rayyan_csv"], index=False, encoding="utf-8")
        output_registry["search_master_xlsx"] = str(paths["search_master_xlsx"])
        output_registry["search_master_csv"] = str(paths["search_master_csv"])
        output_registry["rayyan_csv"] = str(paths["rayyan_csv"])

        examples = [build_naming_example(r) for r in metadata_records[:40]]
        write_text(paths["naming_examples"], "\n".join(examples))
        output_registry["naming_examples"] = str(paths["naming_examples"])

    if local_df is not None and local_sheets is not None:
        sheets = {"MASTER": local_df}
        sheets.update(local_sheets)
        row_dicts_to_excel(sheets, paths["local_xlsx"])
        output_registry["local_xlsx"] = str(paths["local_xlsx"])

    # master workbook
    master_sheets = {}
    if metadata_records:
        master_sheets.update(records_to_export_dfs(metadata_records))
    if local_df is not None:
        master_sheets["LOCAL_MASTER"] = local_df
    if local_sheets:
        for k, v in local_sheets.items():
            master_sheets[f"LOCAL_{k}"] = v
    download_path = paths["download_manifest_csv"]
    if download_path.exists():
        try:
            master_sheets["DOWNLOADS"] = pd.read_csv(download_path)
        except Exception:
            pass
    candidate_path = paths["candidate_links_csv"]
    if candidate_path.exists():
        try:
            master_sheets["CANDIDATES"] = pd.read_csv(candidate_path)
        except Exception:
            pass
    if master_sheets:
        row_dicts_to_excel(master_sheets, paths["master_tables_xlsx"])
        output_registry["master_tables_xlsx"] = str(paths["master_tables_xlsx"])

    write_json(paths["log_json"], output_registry)
    output_registry["log_json"] = str(paths["log_json"])
    return output_registry


# ============================================================
# CLI / MAIN
# ============================================================

def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="NutEV master pipeline v3")
    parser.add_argument("--project-root", default="NUTEV_MASTER_PROJECT", help="Pasta-raiz do projeto")
    parser.add_argument("--taxonomy", default="", help="Caminho opcional para a taxonomia JSON")
    parser.add_argument("--seed-manifest", default="", help="Caminho opcional para o manifesto de sites oficiais")
    parser.add_argument("--email", default="", help="Email para uso educado em APIs")
    parser.add_argument("--mode", default="all", choices=["scaffold", "queries", "search", "download", "local", "all"], help="Fase a executar")
    parser.add_argument("--retmax", type=int, default=DEFAULT_RETMAX, help="Limite por busca bibliográfica e variante")
    parser.add_argument("--web-retmax", type=int, default=DEFAULT_WEB_MAX, help="Limite por busca web e variante")
    parser.add_argument("--download-max", type=int, default=DEFAULT_DOWNLOAD_MAX, help="Máximo de downloads por camada")
    parser.add_argument("--crawl-pages", type=int, default=DEFAULT_CRAWL_PAGES, help="Número máximo de páginas web para rastrear")
    parser.add_argument("--crawl-links-per-page", type=int, default=DEFAULT_CRAWL_LINKS_PER_PAGE, help="Máximo de links candidatos por página")
    parser.add_argument("--no-web-search", action="store_true", help="Desativa busca web geral")
    parser.add_argument("--no-google", action="store_true", help="Desativa Google CSE mesmo se houver credenciais")
    parser.add_argument("--no-download", action="store_true", help="Desativa downloads automáticos")
    parser.add_argument("--no-official-crawl", action="store_true", help="Desativa crawler de fontes oficiais")
    parser.add_argument("--ocr-force", action="store_true", help="Força OCR em PDFs e imagens")
    parser.add_argument("--google-api-key", default=os.environ.get("GOOGLE_API_KEY", ""), help="Google API key (opcional)")
    parser.add_argument("--google-cse-id", default=os.environ.get("GOOGLE_CSE_ID", ""), help="Google CSE ID / cx (opcional)")
    return parser.parse_args()

def main() -> None:
    args = parse_args()
    project_root = Path(args.project_root).resolve()
    scaffold_project(project_root)
    paths = build_default_paths(project_root)

    script_dir = Path(__file__).resolve().parent
    taxonomy_path = Path(args.taxonomy).resolve() if args.taxonomy else paths["taxonomy"]
    if not taxonomy_path.exists() and (script_dir / "nutev_keyword_taxonomy_v3.json").exists():
        taxonomy_path = script_dir / "nutev_keyword_taxonomy_v3.json"
    if not taxonomy_path.exists():
        raise FileNotFoundError(f"Taxonomia não encontrada em {taxonomy_path}")

    seed_manifest_path = Path(args.seed_manifest).resolve() if args.seed_manifest else paths["sources"]
    if not seed_manifest_path.exists() and (script_dir / "nutev_official_sources_manifest_v3.json").exists():
        seed_manifest_path = script_dir / "nutev_official_sources_manifest_v3.json"

    taxonomy = load_taxonomy(taxonomy_path)
    seed_manifest = load_seed_manifest(seed_manifest_path)

    # copia configs para a pasta do projeto
    if not paths["taxonomy"].exists():
        write_json(paths["taxonomy"], taxonomy)
    if seed_manifest and not paths["sources"].exists():
        write_json(paths["sources"], seed_manifest)
    write_text(paths["keywords_md"], render_taxonomy_markdown(taxonomy))

    query_pack = build_query_pack(taxonomy)
    write_query_pack_files(query_pack, project_root)
    log("Pacote de buscas gerado.")

    if not args.email:
        log("Aviso: use --email para operar de forma educada com PubMed/Crossref/Unpaywall.")

    metadata_records: List[SearchRecord] = []
    local_df = None
    local_sheets = None
    official_df = pd.DataFrame()
    candidate_df = pd.DataFrame()

    if args.mode in ["search", "all", "download"]:
        metadata_records, official_df = run_search_layers(
            project_root=project_root,
            taxonomy=taxonomy,
            query_pack=query_pack,
            email=args.email or "your_email@example.org",
            retmax=max(1, args.retmax),
            web_retmax=max(1, args.web_retmax),
            enable_web=not args.no_web_search,
            enable_google=not args.no_google,
            enable_official_crawl=not args.no_official_crawl,
            google_api_key=args.google_api_key,
            google_cse_id=args.google_cse_id,
            seed_manifest=seed_manifest,
        )
        log(f"Registros totais após deduplicação: {len(metadata_records)}")

        if not args.no_download:
            dl_session = session_factory(email=args.email or "your_email@example.org")

            # OA bibliográfico
            dl_oa_df = download_oa_records(dl_session, metadata_records, project_root, max_downloads=max(1, args.download_max))
            log(f"OA bibliográfico processado: {len(dl_oa_df)}")

            # Candidatos a partir das páginas de resultados web
            web_records = [r for r in metadata_records if r.record_kind == "web" and r.prisma_decision != "EXCLUDE"]
            candidate_df = crawl_result_pages_and_collect_candidates(
                dl_session,
                web_records,
                project_root,
                max_pages=max(1, args.crawl_pages),
                max_links_per_page=max(1, args.crawl_links_per_page),
            )
            log(f"Candidatos a partir de resultados web: {len(candidate_df)}")

            # downloads de links oriundos da web
            dl_web_df = download_candidate_links(dl_session, candidate_df, project_root, bucket_name="web_direct_docs", max_downloads=max(1, args.download_max))
            log(f"Links web processados para download: {len(dl_web_df)}")

            # downloads de fontes oficiais
            if official_df is not None and not official_df.empty:
                dl_official_df = download_candidate_links(dl_session, official_df, project_root, bucket_name="official_seed_docs", max_downloads=max(1, args.download_max))
                log(f"Links oficiais processados para download: {len(dl_official_df)}")

    if args.mode in ["local", "all", "download"]:
        local_df, local_sheets = analyze_local_documents(project_root, taxonomy, ocr_force=args.ocr_force)
        log(f"Documentos locais analisados: {len(local_df)}")

    outputs = write_outputs(project_root, metadata_records, local_df, local_sheets)
    log("Pipeline finalizado.")
    for key, path in outputs.items():
        log(f"{key}: {path}")

if __name__ == "__main__":
    main()
